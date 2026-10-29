"""华泰证券对账单解析器（纯函数、无 I/O 直读）。

输入 bytes + file_name，输出结构化 `ParsedBill`。之所以做成纯函数：解析结果要能被
预览（不落库）与提交（落库）两个阶段复用，并可被属性测试反复回放。
所有金额/数量/价格都经 `Decimal(str) → 标度整数`，绝不出现 float。

金额口径（全仓统一，对账直接使用）：`ParsedRow.amount_units` 一律是**带符号的发生金额
（净现金流，已含费用）**——买入为负、卖出为正；非交易资金业务按资金方向（转入为卖出方向
正、转出为买入方向负）；方向不可判定且没有金额流时为 0。账单里原生的「发生金额/发生额」
列已满足该口径，直接采用；「成交金额」毛额（及按量×价推导出的毛额）则按方向换算：
买入 `-(毛额+费用)`、卖出 `毛额-费用`，费用 = 佣金 + 印花税 + 其他费。
"""

from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import xlrd  # type: ignore[import-untyped]  # 第三方 xlrd 无类型桩（types-xlrd 未安装）
from openpyxl import load_workbook  # type: ignore[import-untyped]  # 第三方 openpyxl 无类型桩
from stock_contracts.clock import Clock
from stock_contracts.constants import FILE_MAX_BYTES
from stock_contracts.enums import QuarantineReason, Side, WarningCode
from stock_contracts.errors import ApiError, ErrorCode, FieldError
from stock_core.values import Money, Price, Qty
from stock_platform.symbols import is_convertible_bond

from stock_ledger.parsers.columns import (
    ColumnMapping,
    detect_date_format,
    is_non_trade_summary,
    is_signed_amount_header,
    map_columns,
    normalize_cell,
    normalize_header,
    normalize_summary,
    parse_date_text,
    resolve_side,
)
from stock_ledger.parsers.dedupe import core_signature, dedupe_key

SUPPORTED_SUFFIXES: frozenset[str] = frozenset({".xlsx", ".xls", ".csv"})
_HEADER_SCAN_ROWS = 30
_CURRENCY_NOISE = ("¥", "￥", "元", " ", ",", "\u3000")


@dataclass(slots=True)
class QuarantineCandidate:
    """无法进入明细的行：保留原始内容 + 原因码，供隔离区回填。"""

    line_no: int
    raw: dict[str, str]
    reason_code: str
    suggested_rule_id: str | None = None


@dataclass(slots=True)
class ParsedRow:
    """一条可入明细的交易（金额/数量/价格均为标度整数）。"""

    line_no: int
    trade_date: date
    security_code: str
    security_name: str
    side: str
    quantity_units: int
    price_units: int
    # 带符号发生金额（净现金流，含费用）：买入为负、卖出为正；对账直接用 prev+amount。
    amount_units: int
    commission_units: int
    stamp_tax_units: int
    other_fee_units: int
    cash_balance_units: int | None
    cash_balance_available: bool
    summary: str | None
    summary_code: str | None
    account_masked: str | None
    occurrence_idx: int
    dedupe_key: str
    raw: dict[str, str]
    # 货币基金余额（该行发生后该标的的份额余额，金额 scale=2）。None = 非货基行。
    # 货基单位净值恒 1.0000，没有可估值净值，资产口径按余额计（设计 §5.2）。
    money_fund_balance_units: int | None = None


# 货币基金识别：券商全称含「货币」（归一化折行/空白后），fixtures 3/3 命中。
_MONEY_FUND_KEYWORD = "货币"
# 货基单位净值恒 1.0000；价格列缺失时按此兜底（不伪造份额）。
_ONE_PRICE_UNITS = Price.from_decimal(Decimal(1)).units


def _is_money_fund_name(name: str) -> bool:
    return _MONEY_FUND_KEYWORD in "".join(name.split())


def _is_money_fund_row(row: ParsedRow) -> bool:
    """货基行识别：名称含「货币」，或「买/卖价格 = 1 且佣金/印花税/其他费全 0」。

    后者是对账口径的兜底（部分券商货基全称不含「货币」）；`quantity_units != 0`
    排除转托管确认等零份额噪音行，避免整只标的被误判成货基。
    """
    if _is_money_fund_name(row.security_name):
        return True
    return (
        row.side in (Side.BUY.value, Side.SELL.value)
        and row.price_units == _ONE_PRICE_UNITS
        and row.quantity_units != 0
        and row.commission_units == 0
        and row.stamp_tax_units == 0
        and row.other_fee_units == 0
    )


def _money_fund_codes(rows: list[ParsedRow]) -> set[str]:
    """账单中所有货基标的代码（名称 + 价格 1 零费用两条口径的并集）。"""
    return {row.security_code for row in rows if _is_money_fund_row(row)}


# 零金额信息行（纯通知，无资金意义）：申购中签/新股入账/托管转出/托管转入。
_ZERO_AMOUNT_INFO_KEYWORDS: tuple[str, ...] = ("申购中签", "新股入账", "托管转出", "托管转入")
# 可转债：数量单位是手（1 手 = 10 张），成交/行情价格是单张金额。
_CBOND_LOT_SIZE = 10
# 「qty × 10 × 单张价 ≈ |amount|」容差（amount 含费用，对齐估值反推的 2% 口径）。
_CBOND_AMOUNT_TOLERANCE_PERCENT = 2


def _is_zero_amount_info_summary(summary: str) -> bool:
    """摘要归一化后命中零金额信息关键词（申购中签/新股入账/托管转出入）。"""
    if not summary:
        return False
    normalized = normalize_summary(summary)
    return any(keyword in normalized for keyword in _ZERO_AMOUNT_INFO_KEYWORDS)


def _is_cbond_subscription_name(name: str) -> bool:
    """可转债申购（发债）行：名称去空白后以「发债」结尾（如「曙26发债」）。"""
    return "".join(name.split()).endswith("发债")


def _is_cbond_amount_inconsistent(row: ParsedRow) -> bool:
    """可转债行「数量(手) × 10 × 单张价」与 |发生金额| 偏差超容差（留痕不阻断）。

    只校验有明确买卖方向、量价额俱全的行；申购代码（754/780 等名称以发债结尾）
    与上市转债代码（is_convertible_bond）都按同一口径校验。
    """
    if (
        row.side not in (Side.BUY.value, Side.SELL.value)
        or row.quantity_units <= 0
        or row.price_units <= 0
        or row.amount_units == 0
    ):
        return False
    if not (
        is_convertible_bond(row.security_code) or _is_cbond_subscription_name(row.security_name)
    ):
        return False
    expected = Money.from_decimal(
        Qty(row.quantity_units).to_decimal() * _CBOND_LOT_SIZE * Price(row.price_units).to_decimal()
    ).units
    if expected <= 0:
        return False
    return abs(expected - abs(row.amount_units)) * 100 > _CBOND_AMOUNT_TOLERANCE_PERCENT * expected


def _share_value_units(shares_units: int, price_units: int) -> int:
    """份额 × 单位净值 → 金额标度整数（Money scale=2，ROUND_HALF_UP）。"""
    return Money.from_decimal(
        Qty(shares_units).to_decimal() * Price(price_units).to_decimal()
    ).units


def _read_fund_holding_anchors(rows: list[list[str]], codes: set[str]) -> dict[str, int]:
    """解析「基金持仓」区块，返回 `{货基代码: 期末参考市值(分)}`。

    区块在流水「合计」之后（真实对账单实测表头：基金代码/基金名称/基金份额/单位净值/
    参考市值/摊薄成本价/持仓盈亏），无独立日期列，锚点日按账单最大流水日。缺失
    `参考市值` 时用 `份额 × 单位净值` 兜底；非货基代码忽略。
    """
    header_idx: int | None = None
    for idx, row in enumerate(rows):
        normalized = [normalize_header(cell) for cell in row]
        if "基金代码" in normalized and "基金份额" in normalized:
            header_idx = idx
            break
    if header_idx is None:
        return {}
    header = [normalize_header(cell) for cell in rows[header_idx]]
    code_idx = header.index("基金代码")
    share_idx = header.index("基金份额")
    nav_idx = header.index("单位净值") if "单位净值" in header else None
    value_idx = header.index("参考市值") if "参考市值" in header else None
    anchors: dict[str, int] = {}
    for row in rows[header_idx + 1 :]:
        if not any(row):
            continue
        code = normalize_cell(row[code_idx]) if code_idx < len(row) else ""
        if not code:
            continue
        if not code.isdigit():
            # 区块以「合计/声明」等非代码行收尾。
            break
        if code not in codes:
            continue
        value_units: int | None = None
        if value_idx is not None and value_idx < len(row) and row[value_idx]:
            try:
                value_units = _units(row[value_idx], Money)
            except _RowError:
                value_units = None
        if value_units is None and share_idx < len(row) and row[share_idx]:
            try:
                shares = _units(row[share_idx], Qty)
                nav_units = _units(row[nav_idx], Price) if nav_idx is not None else 0
                if nav_units:
                    value_units = _share_value_units(shares, nav_units)
            except _RowError:
                value_units = None
        if value_units is not None:
            anchors[code] = value_units
    return anchors


def _apply_money_fund_balances(rows: list[ParsedRow], anchors: dict[str, int]) -> set[str]:
    """按余额累计（+ 账单「基金持仓」期末锚点反推期初）逐行写货币基金余额。

    全部在**金额标度（分）**上运行：份额增量先按「份额 × 单位净值(1.0000)」换算成金额，
    避免 Qty(scale=4) 与 Money(scale=2) 混算。返回无锚点的货基代码集合（期初按 0
    推演，需在 assumptions 留痕）。
    """
    codes = _money_fund_codes(rows)
    if not codes:
        return set()
    ordered = sorted(rows, key=lambda row: (row.trade_date, row.line_no))
    deltas: dict[str, int] = {}
    for row in ordered:
        if row.security_code not in codes:
            continue
        deltas[row.security_code] = deltas.get(row.security_code, 0) + _money_fund_row_delta(row)
    missing = {code for code in codes if code not in anchors}
    # 期初余额 = 锚点金额 − 账单内净增量；无锚点按 0（假设留痕）。
    opening: dict[str, int] = {}
    for code in codes:
        # 有锚点：期初 = 锚点金额 − 账单内净增量；无锚点：期初按 0 推演（assumptions 留痕）。
        opening[code] = anchors[code] - deltas.get(code, 0) if code in anchors else 0
    running = dict(opening)
    for row in ordered:
        if row.security_code not in codes:
            continue
        running[row.security_code] += _money_fund_row_delta(row)
        row.money_fund_balance_units = running[row.security_code]
    return missing


def _money_fund_row_delta(row: ParsedRow) -> int:
    """货基行的余额增量（金额分）：卖出取负，其余（含红利再投）取正。"""
    shares = -row.quantity_units if row.side == Side.SELL.value else row.quantity_units
    return _share_value_units(shares, row.price_units or _ONE_PRICE_UNITS)


@dataclass(slots=True)
class ParsedBill:
    """一次解析的完整结果（预览与提交共用）。"""

    rows: list[ParsedRow] = field(default_factory=list)
    columns: dict[str, str] = field(default_factory=dict)
    unresolved: list[str] = field(default_factory=list)
    date_format: str | None = None
    warnings: list[WarningCode] = field(default_factory=list)
    quarantined: list[QuarantineCandidate] = field(default_factory=list)
    assumptions: dict[str, Any] = field(default_factory=dict)

    @property
    def row_count(self) -> int:
        return len(self.rows) + len(self.quarantined)


class _RowError(Exception):
    """单行解析失败（内部信号，外部统一转隔离）。"""

    __slots__ = ("reason",)

    def __init__(self, reason: QuarantineReason) -> None:
        super().__init__(reason.value)
        self.reason = reason


@dataclass(slots=True)
class _Candidate:
    parsed: ParsedRow
    raw_values: tuple[str, ...]


def _validate_file(data: bytes, file_name: str) -> str:
    suffix = Path(file_name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ApiError(
            ErrorCode.IMPORT_UNSUPPORTED_FILE,
            detail=f"不支持的扩展名: {suffix or file_name!r}",
        )
    if len(data) > FILE_MAX_BYTES:
        raise ApiError(ErrorCode.IMPORT_FILE_TOO_LARGE, detail="文件超过 20MB")
    return suffix


def _decode_csv(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_rows(data: bytes, suffix: str) -> list[list[str]]:
    """把三种文件格式统一读成字符串二维表（数值单元格经 str，不引入 float）。"""
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [
                [normalize_cell(cell) for cell in row] for row in sheet.iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    if suffix == ".xls":
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        return [
            [normalize_cell(sheet.cell_value(r, c)) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
    reader = csv.reader(io.StringIO(_decode_csv(data)))
    return [[normalize_cell(cell) for cell in row] for row in reader]


def _locate_header(
    rows: list[list[str]], mapping: dict[str, str] | None
) -> tuple[int, list[str], list[list[str]]]:
    """在前若干行里挑“识别到最多列”的一行作表头（账单常有抬头/说明行）。"""
    best_idx = 0
    best_count = 0
    best_headers: list[str] = rows[0] if rows else []
    for idx in range(min(len(rows), _HEADER_SCAN_ROWS)):
        found: ColumnMapping = map_columns(rows[idx], mapping=mapping)
        count = len(found.columns)
        if count > best_count:
            best_idx, best_count, best_headers = idx, count, rows[idx]
    return best_idx, best_headers, rows[best_idx + 1 :]


def _decimal(raw: str) -> Decimal:
    """单元格文本 → Decimal（去千分位/货币符号）；空值按 0。"""
    cleaned = raw
    for noise in _CURRENCY_NOISE:
        cleaned = cleaned.replace(noise, "")
    if cleaned in ("", "-", "+"):
        return Decimal(0)
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise _RowError(QuarantineReason.INVALID_ROW) from exc


def _units(raw: str, kind: type[Money] | type[Price] | type[Qty]) -> int:
    return kind.from_decimal(_decimal(raw)).units


def _gross_to_net_units(gross_units: int, side: str, fees_units: int) -> int:
    """毛成交额 → 带符号净现金流（分）：买入 `-(毛额+费用)`，卖出 `毛额-费用`。"""
    if side == Side.SELL.value:
        return gross_units - fees_units
    return -(gross_units + fees_units)


def _net_from_gross_units(gross_units: int, side: str, fees_units: int) -> tuple[int, bool]:
    """毛额 → 净现金流；方向未知时无法定号（置 0 并回报），否则按买卖方向换算。"""
    if side == Side.UNKNOWN.value:
        return 0, gross_units != 0
    return _gross_to_net_units(gross_units, side, fees_units), False


def _net_amount_units(
    raw_text: str, *, signed: bool, side: str, fees_units: int
) -> tuple[int, bool]:
    """金额列原文 → 带符号净现金流；返回 `(units, 是否因方向未知而无法定号)`。

    - 「发生金额/发生额」原生带符号 → 直接采用（费用已含在列内，绝不重复扣减）；
    - 毛额列里出现负数 → 只可能是已带符号的净额（毛额不可能为负），按净额采用；
    - 其余按方向换算；方向 unknown 时资金流向不可判定，置 0 并回报调用方留痕。
    """
    raw_units = _units(raw_text, Money)
    if signed or raw_units < 0:
        return raw_units, False
    return _net_from_gross_units(raw_units, side, fees_units)


def _build_candidate(
    raw: dict[str, str],
    field_headers: dict[str, str],
    date_format_code: str,
    line_no: int,
    *,
    amount_signed: bool,
) -> tuple[ParsedRow, bool, bool]:
    """单行 → ParsedRow；失败抛 _RowError 由调用方隔离。

    返回 `(行, 金额是否按量×价派生, 金额是否因方向未知而无法定号)`。
    """

    def value(name: str) -> str:
        header = field_headers.get(name)
        return raw.get(header, "") if header is not None else ""

    code = value("security_code")
    if not code:
        # 非交易类业务（银行转存/转取、股息入帐等）证券代码可空，用 CASH 占位导入。
        if is_non_trade_summary(value("summary")):
            code = "CASH"
        else:
            raise _RowError(QuarantineReason.INVALID_ROW)
    # 方向：买卖标志列可缺失/为空，此时回退摘要识别，仍无则 unknown（不隔离该行）。
    side = resolve_side(value("side"), value("summary"))
    try:
        year, month, day = parse_date_text(value("trade_date"), date_format_code)
        trade_date = date(year, month, day)
    except ValueError as exc:
        raise _RowError(QuarantineReason.INVALID_ROW) from exc

    quantity_units = _units(value("quantity"), Qty)
    price_units = _units(value("price"), Price)

    commission_units = _units(value("commission"), Money)
    stamp_tax_units = _units(value("stamp_tax"), Money)
    other_fee_units = _units(value("other_fee"), Money)
    fees_units = commission_units + stamp_tax_units + other_fee_units

    amount_text = value("amount")
    if amount_text:
        amount_units, amount_unsigned = _net_amount_units(
            amount_text, signed=amount_signed, side=side, fees_units=fees_units
        )
        amount_derived = False
    else:
        # 无金额列：先按量×价推导毛额，再按方向换算成净现金流（不能用原始毛额入库）。
        gross_units = Money.from_decimal(
            Qty(quantity_units).to_decimal() * Price(price_units).to_decimal()
        ).units
        amount_units, amount_unsigned = _net_from_gross_units(gross_units, side, fees_units)
        amount_derived = True

    balance_text = value("cash_balance")
    cash_balance_units: int | None = None
    if balance_text:
        cash_balance_units = _units(balance_text, Money)

    # 「申购中签/新股入账/托管转出」等纯通知行：无发生金额且无资金余额 → 隔离为
    # ZERO_AMOUNT_INFO，不进明细（否则会形成无意义持仓/回合噪音，如中签通知行使
    # 持仓形成日早于真实缴款日）。同类摘要但有金额的行（真实缴款）不受影响。
    if (
        amount_units == 0
        and cash_balance_units is None
        and _is_zero_amount_info_summary(value("summary"))
    ):
        raise _RowError(QuarantineReason.ZERO_AMOUNT_INFO)

    # 五个金额字段（佣金/印花税/其他费/发生金额/资金余额）全为 0：没有资金流语义
    # （典型是货基份额再投、转托管确认等噪音行），不入明细，隔离为 INVALID_AMOUNT。
    # 资金余额列缺失（None）不算 0，交后续对账按 MISSING_CASH_BALANCE 降级处理。
    if (
        commission_units == 0
        and stamp_tax_units == 0
        and other_fee_units == 0
        and amount_units == 0
        and cash_balance_units == 0
    ):
        raise _RowError(QuarantineReason.INVALID_AMOUNT)

    summary = value("summary") or None
    # 账号明文已在 parse_bill 入口统一抹成 MD5（空账号留空串 → 这里转 None）。
    # 此值同时作为 ParsedRow.account_masked（落库列）与下游去重键的业务要素。
    account_masked = value("account") or None
    # 非交易行证券名称为空时，用摘要填充（如「银行转存」），避免默认成 CASH 无法辨识业务。
    security_name = value("security_name") or (summary or "资金业务" if code == "CASH" else code)

    row = ParsedRow(
        line_no=line_no,
        trade_date=trade_date,
        security_code=code,
        security_name=security_name,
        side=side,
        quantity_units=quantity_units,
        price_units=price_units,
        amount_units=amount_units,
        commission_units=commission_units,
        stamp_tax_units=stamp_tax_units,
        other_fee_units=other_fee_units,
        cash_balance_units=cash_balance_units,
        cash_balance_available=cash_balance_units is not None,
        summary=summary,
        summary_code=None,
        account_masked=account_masked,
        occurrence_idx=0,
        dedupe_key="",
        raw=raw,
    )
    return row, amount_derived, amount_unsigned


def parse_bill(
    data: bytes,
    file_name: str,
    *,
    mapping: dict[str, str] | None,
    date_format: str | None,
    clock: Clock,
) -> ParsedBill:
    """解析华泰账单；所有业务失败以 ApiError/隔离行表达，不抛裸异常。"""
    suffix = _validate_file(data, file_name)
    raw_rows = _read_rows(data, suffix)
    header_idx, headers, data_rows = _locate_header(raw_rows, mapping)

    found = map_columns(headers, mapping=mapping)
    missing_required = found.missing_required()
    if missing_required:
        raise ApiError(
            ErrorCode.IMPORT_COLUMNS_UNRESOLVED,
            errors=[
                FieldError("columns", "missing_required_column", f"缺少必需列: {name}")
                for name in missing_required
            ],
            detail="账单缺少必需列",
        )
    field_headers = found.field_headers()
    # 金额列语义由原始表头决定：「发生金额/发生额」带符号直接用，其余按毛额+方向换算。
    amount_signed = is_signed_amount_header(field_headers.get("amount"))

    warnings: list[WarningCode] = []
    assumptions: dict[str, Any] = {}

    trade_date_header = field_headers.get("trade_date")
    trade_date_idx = headers.index(trade_date_header) if trade_date_header is not None else None
    samples: list[str] = []
    if trade_date_idx is not None:
        for row in data_rows:
            if trade_date_idx < len(row):
                cell = row[trade_date_idx]
                # 「合计」之后不是交易行：样本遇到它即停，避免把「合计」当日期样本。
                if normalize_header(cell) == "合计":
                    break
                if cell:
                    samples.append(cell)
            if len(samples) >= 20:
                break

    detected_format, detected_warnings = detect_date_format(samples)
    if date_format is not None:
        effective_format = date_format
    else:
        effective_format = detected_format or "YYYY-MM-DD"
        warnings.extend(detected_warnings)
        if (
            detected_format is None and samples
        ) or WarningCode.DATE_FORMAT_AMBIGUOUS in detected_warnings:
            assumptions["date_format"] = effective_format

    candidates: list[_Candidate] = []
    quarantined: list[QuarantineCandidate] = []
    derived_count = 0
    unsigned_amount_count = 0
    # 流水明细以「合计」行收尾；其后的持仓区块用于货币基金余额锚点（见 §5.3）。
    stop_row_index: int | None = None

    for offset, row in enumerate(data_rows):
        line_no = header_idx + offset + 2
        if not any(cell for cell in row):
            continue
        # 流水明细以「合计」行收尾：其后是股票/基金持仓、声明等非交易区块。
        # 不在这里停止的话，上百行持仓会被当成坏行写进隔离区，污染“隔离数据”。
        if (
            trade_date_idx is not None
            and trade_date_idx < len(row)
            and normalize_header(row[trade_date_idx]) == "合计"
        ):
            stop_row_index = header_idx + offset + 1
            break
        raw = {header: (row[i] if i < len(row) else "") for i, header in enumerate(headers)}
        # 把账号列明文替换成 MD5 摘要：隔离区 raw 和预览 columns_detail.sample
        # 都从这里拿原始行字典，源头抹一次即可，下游都安全。
        account_header = field_headers.get("account")
        if account_header is not None:
            account_text = raw.get(account_header, "")
            raw[account_header] = (
                hashlib.md5(account_text.encode("utf-8")).hexdigest() if account_text else ""
            )
        try:
            parsed, amount_derived, amount_unsigned = _build_candidate(
                raw,
                field_headers,
                effective_format,
                line_no,
                amount_signed=amount_signed,
            )
        except _RowError as err:
            quarantined.append(
                QuarantineCandidate(line_no=line_no, raw=raw, reason_code=err.reason.value)
            )
            continue
        if amount_derived:
            derived_count += 1
        if amount_unsigned:
            unsigned_amount_count += 1
        candidates.append(
            _Candidate(parsed=parsed, raw_values=tuple(raw.get(h, "") for h in headers))
        )

    kept: list[_Candidate] = []
    seen_full: set[tuple[str, ...]] = set()
    for cand in candidates:
        if cand.raw_values in seen_full:
            quarantined.append(
                QuarantineCandidate(
                    line_no=cand.parsed.line_no,
                    raw=cand.parsed.raw,
                    reason_code=QuarantineReason.DUPLICATE_IN_FILE.value,
                )
            )
            continue
        seen_full.add(cand.raw_values)
        kept.append(cand)

    counts: dict[str, int] = {}
    rows: list[ParsedRow] = []
    for cand in kept:
        parsed_row = cand.parsed
        signature = core_signature(parsed_row)
        parsed_row.occurrence_idx = counts.get(signature, 0)
        counts[signature] = parsed_row.occurrence_idx + 1
        parsed_row.dedupe_key = dedupe_key(parsed_row)
        rows.append(parsed_row)

    if "cash_balance" not in field_headers:
        warnings.append(WarningCode.MISSING_CASH_BALANCE)
    # 货币基金余额：按份额累计推演；若账单含「基金持仓」区块，用其参考市值作期末锚点
    # 反推期初（设计 §5.3）。无锚点的代码期初按 0，写 assumptions 留痕，不伪造余额。
    fund_rows = raw_rows[stop_row_index + 1 :] if stop_row_index is not None else []
    fund_anchors = _read_fund_holding_anchors(fund_rows, _money_fund_codes(rows))
    unanchored = _apply_money_fund_balances(rows, fund_anchors)
    if unanchored:
        assumptions["money_fund_opening_zero"] = sorted(unanchored)
    if rows and rows[0].side != "buy":
        warnings.append(WarningCode.FIRST_ROW_NOT_BUY)
    if any(q.reason_code == QuarantineReason.DUPLICATE_IN_FILE.value for q in quarantined):
        warnings.append(WarningCode.DUPLICATE_ROWS)
    if derived_count:
        assumptions["amount_derived_rows"] = derived_count
    if unsigned_amount_count:
        # 毛额列 + 方向 unknown：资金流向不可判定，金额保守置 0，不伪造正负号。
        assumptions["amount_unsigned_rows"] = unsigned_amount_count
    # 可转债口径留痕：数量单位是手（1 手 = 10 张）、价格是单张金额，「qty × 10 ×
    # 单张价」应与 |发生金额| 大致相等；偏差超容差说明券商口径异常，WARN 不阻断。
    cbond_inconsistent = sum(1 for row in rows if _is_cbond_amount_inconsistent(row))
    if cbond_inconsistent:
        warnings.append(WarningCode.CBOND_AMOUNT_INCONSISTENT)
        assumptions["cbond_amount_inconsistent_rows"] = cbond_inconsistent

    return ParsedBill(
        rows=rows,
        columns=found.columns,
        unresolved=found.unresolved,
        date_format=effective_format,
        warnings=warnings,
        quarantined=quarantined,
        assumptions=assumptions,
    )


__all__ = [
    "SUPPORTED_SUFFIXES",
    "ParsedBill",
    "ParsedRow",
    "QuarantineCandidate",
    "parse_bill",
]
